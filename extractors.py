"""
extractors.py
=============================================================================
Responsible for ONE thing only: reading a supported file and returning its
contents as a list of (location_label, text) tuples.

  location_label  – human-readable position, e.g. "Line 12",
                    "Sheet: Sales | Cell C4",  "Page 2 | Line 7"
  text            – the raw string of that unit (line / cell / paragraph)

The scanner (scanner.py) never touches files directly — it only receives
these tuples, so adding a new file type here requires zero changes elsewhere.

Supported formats
-----------------
  .txt   plain text
  .csv   comma/semicolon/tab-separated values
  .docx  Microsoft Word  (requires: pip install python-docx)
  .pdf   PDF documents   (requires: pip install pdfplumber)
  .xlsx  Excel workbooks (requires: pip install openpyxl)

To add a new format:
  1. Write a new extract_<ext>(path: Path) -> list[tuple[str, str]] function.
  2. Register it in the EXTRACTORS dict at the bottom of this file.
  That is all — no changes needed in scanner.py or pii_scanner.py.
=============================================================================
"""

from pathlib import Path

# ── Optional third-party imports (graceful degradation) ──────────────────────
# If a library is missing the corresponding extractor returns a single
# informational tuple instead of crashing the whole application.

try:
    import docx as _docx     # python-docx
    _DOCX_OK = True
except ImportError:
    _DOCX_OK = False

try:
    import pdfplumber         # pdfplumber
    _PDF_OK = True
except ImportError:
    _PDF_OK = False

try:
    import openpyxl           # openpyxl
    _XLSX_OK = True
except ImportError:
    _XLSX_OK = False


# =============================================================================
# INTERNAL HELPERS
# =============================================================================

def _safe_read(path: Path) -> str:
    """
    Read a text file with encoding fallback.

    Attempt order:
      1. utf-8-sig  – UTF-8 with optional BOM (common on Windows)
      2. latin-1    – decodes every byte 0x00-0xFF without error,
                      covers most legacy Windows-created text files.

    Returns the file's content as a plain string.
    """
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    # Absolute last resort — replace undecodable bytes with a placeholder
    return path.read_text(encoding="utf-8", errors="replace")


def _col_letter(zero_based_index: int) -> str:
    """
    Convert a 0-based column index to an Excel-style letter.
      0 → A,  25 → Z,  26 → AA,  27 → AB, …

    Used to produce human-friendly location labels like "Col B" or "Col AA".
    """
    result = ""
    n = zero_based_index
    while True:
        result = chr(65 + n % 26) + result
        n = n // 26 - 1
        if n < 0:
            break
    return result


# =============================================================================
# EXTRACTORS — one function per file format
# =============================================================================

def extract_txt(path: Path) -> list[tuple[str, str]]:
    """
    Plain text files (.txt).

    Strategy: read the whole file, split on any line ending.

    Edge cases handled:
      • Mixed line endings (\\r\\n, \\n, \\r) — str.splitlines() handles all.
      • BOM at file start — stripped by utf-8-sig encoding.
      • Completely empty file — returns [] (no crash).
      • Non-UTF-8 characters — handled by _safe_read fallback chain.

    Returns one tuple per non-empty line:
      ("Line 12", "some text on line 12")
    """
    text = _safe_read(path)

    units = []
    for i, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        if stripped:                          # skip blank lines
            units.append((f"Line {i}", stripped))
    return units


def extract_csv(path: Path) -> list[tuple[str, str]]:
    """
    CSV / TSV / delimiter-separated files (.csv).

    Strategy: use csv.Sniffer to auto-detect the delimiter, then
    iterate every cell individually so PII within a cell is found
    regardless of its position in the row.

    Edge cases handled:
      • Different delimiters (, ; \\t |) — Sniffer detects automatically.
      • Quoted fields with embedded commas or newlines — csv.reader handles.
      • Completely empty cells — skipped silently.
      • Header row — treated like any other row (may itself contain PII).
      • Encoding — same utf-8-sig → latin-1 fallback as .txt.

    Returns one tuple per non-empty cell:
      ("Row 3 | Col B", "cell value")
    """
    import csv
    import io

    raw = _safe_read(path)

    # Detect delimiter from first 2 KB; fall back to standard CSV if it fails
    try:
        dialect = csv.Sniffer().sniff(raw[:2048], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    units = []
    reader = csv.reader(io.StringIO(raw), dialect)
    for row_idx, row in enumerate(reader, start=1):
        for col_idx, cell in enumerate(row):
            text = cell.strip()
            if text:
                col_letter = _col_letter(col_idx)
                units.append((f"Row {row_idx} | Col {col_letter}", text))
    return units


def extract_docx(path: Path) -> list[tuple[str, str]]:
    """
    Microsoft Word documents (.docx).

    Strategy: extract text from three areas of the document:
      1. Body paragraphs — the main document text.
      2. Tables          — each cell scanned individually.
      3. Headers/footers — often contain names, dates, org info.

    Edge cases handled:
      • Password-protected / corrupt files — caught with a generic
        Exception handler; returns a single error tuple.
      • Empty paragraphs / cells — skipped silently.
      • Missing python-docx library — returns an info tuple explaining
        the library must be installed.

    Returns tuples like:
      ("Paragraph 3",           "some paragraph text")
      ("Table 1 | Row 2 | Col 3", "cell content")
      ("Section 1 Header",      "header text")
    """
    if not _DOCX_OK:
        return [("ERROR", "python-docx is not installed. Run: pip install python-docx")]

    units = []
    try:
        doc = _docx.Document(str(path))

        # ── Body paragraphs ──
        for i, para in enumerate(doc.paragraphs, start=1):
            text = para.text.strip()
            if text:
                units.append((f"Paragraph {i}", text))

        # ── Tables ──
        # doc.tables gives all top-level tables; nested tables are ignored
        # (rare in practice and hard to address without recursion).
        for t_idx, table in enumerate(doc.tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                for c_idx, cell in enumerate(row.cells, start=1):
                    text = cell.text.strip()
                    if text:
                        units.append((
                            f"Table {t_idx} | Row {r_idx} | Col {c_idx}",
                            text,
                        ))

        # ── Headers and footers ──
        for sec_idx, section in enumerate(doc.sections, start=1):
            for part_name, part in [("Header", section.header),
                                     ("Footer", section.footer)]:
                if part is None:
                    continue
                for para in part.paragraphs:
                    text = para.text.strip()
                    if text:
                        units.append((f"Section {sec_idx} {part_name}", text))

    except Exception as exc:
        units.append(("ERROR", f"Could not read .docx file: {exc}"))

    return units


def extract_pdf(path: Path) -> list[tuple[str, str]]:
    """
    PDF documents (.pdf).

    Strategy: use pdfplumber to extract text page-by-page, then split
    each page's text into lines.

    Edge cases handled:
      • Scanned / image-only PDFs — pdfplumber returns None or empty string
        for image pages; these are flagged with an informational message
        rather than silently dropped.
      • Encrypted PDFs — caught; returns an error tuple.
      • Multi-column layouts — pdfplumber preserves approximate reading order.
      • Garbled text from bad PDF fonts — best-effort; no crash.
      • Missing pdfplumber library — returns an info tuple.

    Returns tuples like:
      ("Page 2 | Line 7", "text on that line")
      ("Page 4",          "[no extractable text — possibly a scanned image]")
    """
    if not _PDF_OK:
        return [("ERROR", "pdfplumber is not installed. Run: pip install pdfplumber")]

    units = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            if not pdf.pages:
                return [("INFO", "PDF contains no pages")]

            for pg_idx, page in enumerate(pdf.pages, start=1):
                text = page.extract_text()

                # Image-only page — flag it but continue with remaining pages
                if not text or not text.strip():
                    units.append((
                        f"Page {pg_idx}",
                        "[no extractable text — possibly a scanned image]",
                    ))
                    continue

                for ln_idx, line in enumerate(text.splitlines(), start=1):
                    line = line.strip()
                    if line:
                        units.append((f"Page {pg_idx} | Line {ln_idx}", line))

    except Exception as exc:
        units.append(("ERROR", f"Could not read .pdf file: {exc}"))

    return units


def extract_xlsx(path: Path) -> list[tuple[str, str]]:
    """
    Microsoft Excel workbooks (.xlsx).

    Strategy: iterate every worksheet, every row, every cell.

    Edge cases handled:
      • Multiple worksheets — all are scanned; sheet name is in the label.
      • Merged cells — openpyxl returns the value only in the top-left cell;
        the remaining merged cells read as None and are skipped automatically.
      • Formula cells — data_only=True reads the cached result stored by Excel
        the last time the file was saved (not a live recalculation).
      • Date-formatted cells — openpyxl returns datetime objects; we convert
        to string with str() so the scanner receives plain text.
      • Empty sheets — skipped silently (no rows yield no tuples).
      • Password-protected workbooks — caught; returns an error tuple.
      • Missing openpyxl library — returns an info tuple.

    Returns tuples like:
      ("Sheet: Sales | Cell C14", "1234.56")
      ("Sheet: Contacts | Cell A2", "John Smith")
    """
    if not _XLSX_OK:
        return [("ERROR", "openpyxl is not installed. Run: pip install openpyxl")]

    units = []
    try:
        # data_only=True  → read cached formula values, not formula strings
        wb = openpyxl.load_workbook(str(path), data_only=True)

        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    text = str(cell.value).strip()
                    if text:
                        label = f"Sheet: {sheet.title} | Cell {cell.column_letter}{cell.row}"
                        units.append((label, text))

    except Exception as exc:
        units.append(("ERROR", f"Could not read .xlsx file: {exc}"))

    return units


# =============================================================================
# EXTRACTOR DISPATCH TABLE
# =============================================================================
# Maps a lowercase file extension to its extractor function.
# This is the ONLY place that needs to change when adding a new file type:
#   1. Write extract_<ext>() above.
#   2. Add one line here.

EXTRACTORS: dict[str, callable] = {
    ".txt":  extract_txt,
    ".csv":  extract_csv,
    ".docx": extract_docx,
    ".pdf":  extract_pdf,
    ".xlsx": extract_xlsx,
}

# Convenience set used by scanner.py for quick extension checks
SUPPORTED_EXTENSIONS: set[str] = set(EXTRACTORS.keys())
